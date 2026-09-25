import os
import re
import sys
import time
import argparse
from datetime import datetime

import requests
import pandas as pd
from openpyxl.styles import Font, PatternFill

API = "https://api.github.com"
EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
# domains where a "firstname@domain" guess makes no sense
SKIP = ("github.io", "github.com", "linkedin.com", "twitter.com", "x.com", "youtube.com",
        "medium.com", "instagram.com", "facebook.com", "t.me", "linktr.ee", "gmail.com",
        "hashnode.dev", "dev.to", "substack.com", "blogspot.com", "wordpress.com", "vercel.app",
        "netlify.app", "about.me", "bit.ly", "google.com", "gitlab.com")


def session():
    s = requests.Session()
    s.headers["Accept"] = "application/vnd.github+json"
    tok = os.getenv("GH_TOKEN")
    if tok:
        s.headers["Authorization"] = f"Bearer {tok}"
    return s


def get(s, url, **params):
    r = s.get(url, params=params, timeout=15)
    if r.status_code == 403 and r.headers.get("x-ratelimit-remaining") == "0":
        wait = int(r.headers.get("x-ratelimit-reset", time.time() + 60)) - int(time.time())
        print(f"rate limited, waiting {wait}s...")
        time.sleep(max(wait, 1) + 1)
        return get(s, url, **params)
    r.raise_for_status()
    return r.json()


def fetch(q, n):
    s = session()
    users = []
    page = 1
    while len(users) < n:
        res = get(s, f"{API}/search/users", q=q, per_page=min(n, 100), page=page)
        items = res.get("items", [])
        if not items:
            break
        users += items
        page += 1
    users = users[:n]

    rows = []
    for i, u in enumerate(users, 1):
        d = get(s, u["url"])
        socials = []
        if "Authorization" in s.headers:  # extra call per user, only worth it with a token
            try:
                socials = [x["url"] for x in get(s, f"{u['url']}/social_accounts")]
            except requests.HTTPError:
                pass
        rows.append({
            "login": d["login"],
            "name": d.get("name"),
            "email": d.get("email"),
            "bio": d.get("bio"),
            "company": d.get("company"),
            "website": d.get("blog"),
            "twitter": d.get("twitter_username"),
            "socials": " ".join(socials),
            "location": d.get("location"),
            "followers": d.get("followers"),
            "repos": d.get("public_repos"),
            "github": d["html_url"],
        })
        print(f"[{i}/{len(users)}] {d['login']}")
    return pd.DataFrame(rows)


def fix_url(u):
    if not isinstance(u, str) or not u.strip():
        return None
    u = u.strip().rstrip("/")
    return u if u.startswith("http") else "https://" + u


def domain(u):
    if not isinstance(u, str):
        return None
    d = re.sub(r"^https?://", "", u).split("/")[0].lower()
    d = d[4:] if d.startswith("www.") else d
    return None if any(d.endswith(x) for x in SKIP) else d


def guess(name, dom):
    # common company email patterns: first@, first.last@, flast@
    if not isinstance(name, str) or not dom:
        return None
    parts = [p for p in re.sub(r"[^a-z0-9 ]", "", name.lower()).split() if len(p) > 1]
    if not parts:
        return None
    f, l = parts[0], parts[-1]
    out = [f"{f}@{dom}"]
    if len(parts) > 1:
        out += [f"{f}.{l}@{dom}", f"{f[0]}{l}@{dom}"]
    return ", ".join(out)


def linkedin(row):
    txt = " ".join(str(row[c]) for c in ("website", "socials", "bio") if pd.notna(row[c]))
    m = re.search(r"(https?://)?(www\.)?linkedin\.com/in/[\w-]+", txt)
    return fix_url(m.group(0)) if m else None


def clean(df):
    df = df.copy()
    for c in df.select_dtypes(["object", "string"]):
        df[c] = df[c].str.strip().replace("", None)

    # pick up emails people put in their bio
    bio_mail = df["bio"].str.extract(f"({EMAIL_RE.pattern})", expand=False)
    df["email"] = df["email"].fillna(bio_mail).str.lower()
    df.loc[~df["email"].fillna("").str.fullmatch(EMAIL_RE.pattern), "email"] = None

    df["website"] = df["website"].map(fix_url)
    df["linkedin"] = df.apply(linkedin, axis=1)
    df["name"] = df["name"].fillna(df["login"]).str.replace(r"[^\w\s.'-]", "", regex=True).str.strip()
    df["name"] = df["name"].map(lambda n: n.title() if n.isupper() else n)
    df["company"] = df["company"].str.lstrip("@")
    df["location"] = df["location"].str.title()
    df["email_guess"] = [guess(n, domain(w)) for n, w in zip(df["name"], df["website"])]

    before = len(df)
    df = df.drop_duplicates("login")
    df = df[~(df["email"].notna() & df.duplicated("email"))]
    print(f"removed {before - len(df)} duplicates")

    df = df.sort_values("followers", ascending=False)
    cols = ["name", "email", "email_guess", "website", "linkedin", "github", "location",
            "company", "followers", "repos", "twitter", "bio"]
    df = df[cols].fillna("N/A")
    df.columns = [c.replace("_", " ").title() for c in cols]
    return df.reset_index(drop=True)


def save(df, path, q):
    stats = pd.DataFrame({
        "Metric": ["Query", "Run at", "Total leads", "With email", "With email guess",
                   "With website", "With LinkedIn"],
        "Value": [q, datetime.now().strftime("%Y-%m-%d %H:%M"), len(df),
                  (df["Email"] != "N/A").sum(), (df["Email Guess"] != "N/A").sum(),
                  (df["Website"] != "N/A").sum(), (df["Linkedin"] != "N/A").sum()],
    })
    with pd.ExcelWriter(path, engine="openpyxl") as w:
        df.to_excel(w, sheet_name="Leads", index=False)
        stats.to_excel(w, sheet_name="Summary", index=False)
        for ws in w.book.worksheets:
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="305496")
            for col in ws.columns:
                width = max(len(str(c.value or "")) for c in col)
                ws.column_dimensions[col[0].column_letter].width = min(width + 2, 50)
            ws.freeze_panes = "A2"
    print(f"saved {len(df)} leads -> {path}")


def to_gsheet(df, name):
    # needs gspread + a service account key in creds.json
    import gspread
    gc = gspread.service_account("creds.json")
    try:
        sh = gc.open(name)
    except gspread.SpreadsheetNotFound:
        sh = gc.create(name)
    ws = sh.sheet1
    ws.clear()
    ws.update([df.columns.tolist()] + df.astype(str).values.tolist())
    print(f"pushed to google sheet: {sh.url}")


def run(a):
    print(f"\n--- run {datetime.now():%Y-%m-%d %H:%M} | q='{a.q}' ---")
    os.makedirs(os.path.dirname(a.out) or ".", exist_ok=True)
    raw = fetch(a.q, a.n)
    raw.to_csv(os.path.join(os.path.dirname(a.out) or ".", "raw.csv"), index=False)
    df = clean(raw)
    save(df, a.out, a.q)
    if a.sheet:
        try:
            to_gsheet(df, a.sheet)
        except Exception as e:
            print(f"google sheet skipped: {e}")


def main():
    p = argparse.ArgumentParser(description="Collect developer leads from GitHub into Excel")
    p.add_argument("-q", default="location:hyderabad followers:>300 type:user",
                   help="github user search query")
    p.add_argument("-n", type=int, default=30, help="number of leads")
    p.add_argument("-o", "--out", default="output/leads.xlsx")
    p.add_argument("--sheet", help="also push to this google sheet (needs creds.json)")
    p.add_argument("--every", type=float, help="rerun every N hours")
    a = p.parse_args()

    while True:
        try:
            run(a)
        except requests.RequestException as e:
            print(f"run failed: {e}", file=sys.stderr)
            if not a.every:
                sys.exit(1)
        if not a.every:
            break
        print(f"next run in {a.every}h (ctrl+c to stop)")
        time.sleep(a.every * 3600)


if __name__ == "__main__":
    main()
